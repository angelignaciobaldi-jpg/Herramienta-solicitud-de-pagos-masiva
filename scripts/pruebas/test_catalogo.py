"""Catálogo de conceptos de pago y documentos adjuntos.

Dos cosas que SIPP impone y la herramienta tiene que reflejar sin inventar: qué
conceptos existen de verdad, y que una cuenta bancaria nueva no se registra sin
su carátula.
"""

from __future__ import annotations

import os

from openpyxl import load_workbook

from core import conceptos, db, documentos, plantilla_excel
from scripts.pruebas import comun


# --------------------------------------------------------------------------- #
#  Conceptos
# --------------------------------------------------------------------------- #
def probar_catalogo_arranca_vacio():
    """sin importar nada, no hay conceptos"""
    assert not conceptos.hay_catalogo()
    assert conceptos.nombres() == []


def probar_importar_no_duplica():
    """volver a importar detecta lo nuevo sin repetir lo que ya estaba"""
    lista = ["VIGILANCIA", "NO DEDUCIBLE", "PAGO TARJETA CREDITO"]
    primera = conceptos.importar(lista, "Abastecedora")
    assert primera == {"nuevos": 3, "ya_estaban": 0}
    segunda = conceptos.importar(lista, "Abastecedora")
    assert segunda == {"nuevos": 0, "ya_estaban": 3}
    assert len(conceptos.listar()) == 3


def probar_es_una_lista_unica_para_el_grupo():
    """el mismo concepto en otra empresa no crea una entrada nueva"""
    conceptos.importar(["VIGILANCIA"], "Abastecedora")
    conceptos.importar(["VIGILANCIA", "OTRO"], "Aske")
    assert len(conceptos.listar()) == 2


def probar_normaliza_para_no_duplicar():
    """«vigilancia» y «VIGILANCIA » son el mismo concepto"""
    conceptos.guardar("VIGILANCIA", "Abastecedora", conceptos.ORIGEN_SIPP)
    assert conceptos.guardar("vigilancia") is None
    assert conceptos.guardar("Vigilancia ") is None
    assert len(conceptos.listar()) == 1


def probar_el_manual_queda_sin_verificar():
    """un concepto capturado a mano se distingue del importado"""
    # La diferencia importa: uno existe con certeza en SIPP y el otro puede no
    # existir, y si no existe, la solicitud que lo use no se podrá capturar.
    conceptos.importar(["VIGILANCIA"], "Abastecedora")
    manual = conceptos.guardar("CONCEPTO INVENTADO")
    importado = [c for c in conceptos.listar() if c.nombre == "VIGILANCIA"][0]
    assert manual is not None and not manual.verificado
    assert importado.verificado


def probar_la_plantilla_ofrece_el_catalogo():
    """los conceptos del catálogo llegan a la plantilla como desplegable"""
    carpeta = comun.carpeta_temporal()
    sin = os.path.join(carpeta, "sin.xlsx")
    plantilla_excel.generar(sin)
    listas_sin = [c.value for c in
                  load_workbook(sin)[plantilla_excel.HOJA_CATALOGOS][2]]
    assert "conceptos" not in listas_sin, (
        "sin catálogo no debe inventarse una lista")

    conceptos.importar(["VIGILANCIA", "NO DEDUCIBLE"], "Abastecedora")
    con = os.path.join(carpeta, "con.xlsx")
    plantilla_excel.generar(con)
    libro = load_workbook(con)
    hoja_cat = libro[plantilla_excel.HOJA_CATALOGOS]
    listas_con = [c.value for c in hoja_cat[2]]
    assert "conceptos" in listas_con
    columna = listas_con.index("conceptos") + 1
    valores = [hoja_cat.cell(row=r, column=columna).value for r in (3, 4)]
    assert sorted(v for v in valores if v) == ["NO DEDUCIBLE", "VIGILANCIA"]


# --------------------------------------------------------------------------- #
#  Documentos
# --------------------------------------------------------------------------- #
def probar_emparejar_caratula_por_nombre():
    """el archivo se reconoce aunque traiga texto de más"""
    carpeta = comun.carpeta_temporal()
    for nombre in ("CARATULA JUAN PEREZ LOPEZ BBVA.pdf",
                   "María Ñuño Gómez.pdf",
                   "estado de cuenta - LUIS ALBERTO DIAZ MORA - agosto.pdf",
                   "ruido que no corresponde a nadie.pdf"):
        comun.pdf_falso(carpeta, nombre)

    casos = [("Juan Pérez López", True), ("MARIA NUNO GOMEZ", True),
             ("Luis Alberto Díaz Mora", True), ("Juan Pérez", True),
             # Un apellido distinto NO es la misma persona.
             ("Juan Pérez Ramírez", False), ("PEDRO SIN ARCHIVO", False)]
    for nombre, esperado in casos:
        hallado = bool(documentos.buscar_en_carpeta(carpeta, nombre))
        assert hallado == esperado, nombre


def probar_asignar_una_carpeta_al_lote():
    """se asigna por nombre y se dice a quién le falta"""
    carpeta = comun.carpeta_temporal()
    comun.pdf_falso(carpeta, "CARATULA JUAN PEREZ LOPEZ.pdf")
    comun.pdf_falso(carpeta, "CARATULA MARIA RUIZ SOTO.pdf")

    lote = comun.lote()
    partidas = [comun.concepto("VIGILANCIA", 1000.0)]
    juan = comun.solicitud(lote.id, "JUAN PEREZ LOPEZ", partidas=partidas)
    comun.solicitud(lote.id, "MARIA RUIZ SOTO",
                    partidas=[comun.concepto("VIGILANCIA", 2000.0)])
    pedro = comun.solicitud(lote.id, "PEDRO SIN ARCHIVO",
                            partidas=[comun.concepto("VIGILANCIA", 500.0)])

    resultado = documentos.asignar_por_carpeta(
        lote.id, carpeta, documentos.TIPO_CARATULA)
    assert resultado["asignados"] == 2
    assert "PEDRO SIN ARCHIVO" in resultado["sin_archivo"]
    assert "caratula" in documentos.de_solicitud(juan.id)
    assert "caratula" not in documentos.de_solicitud(pedro.id)


def probar_reasignar_no_duplica():
    """una solicitud tiene UNA carátula, no un historial"""
    carpeta = comun.carpeta_temporal()
    comun.pdf_falso(carpeta, "CARATULA ANA LOPEZ.pdf")
    lote = comun.lote()
    ana = comun.solicitud(lote.id, "ANA LOPEZ",
                          partidas=[comun.concepto("VIGILANCIA", 100.0)])
    for _ in range(2):
        documentos.asignar_por_carpeta(lote.id, carpeta,
                                       documentos.TIPO_CARATULA)
    assert len(db.listar_documentos(solicitud_id=ana.id)) == 1


def probar_caratula_y_vobo_conviven():
    """los dos tipos de documento se guardan a la vez"""
    lote = comun.lote()
    ana = comun.solicitud(lote.id, "ANA LOPEZ",
                          partidas=[comun.concepto("VIGILANCIA", 100.0)])
    carpeta = comun.carpeta_temporal()
    documentos.registrar(ana.id, comun.pdf_falso(carpeta, "c.pdf"),
                         documentos.TIPO_CARATULA, lote.id)
    documentos.registrar(ana.id, comun.pdf_falso(carpeta, "v.pdf"),
                         documentos.TIPO_VOBO, lote.id)
    assert set(documentos.de_solicitud(ana.id)) == {"caratula", "vobo"}


def probar_archivo_borrado_del_disco_no_se_ofrece():
    """si el archivo desapareció, el motor no lo recibe"""
    lote = comun.lote()
    ana = comun.solicitud(lote.id, "ANA LOPEZ",
                          partidas=[comun.concepto("VIGILANCIA", 100.0)])
    carpeta = comun.carpeta_temporal()
    ruta = comun.pdf_falso(carpeta, "c.pdf")
    documentos.registrar(ana.id, ruta, documentos.TIPO_CARATULA, lote.id)
    os.remove(ruta)
    assert "caratula" not in documentos.de_solicitud(ana.id)


def probar_aviso_de_caratula_no_depende_del_flag():
    """se avisa por transferencia sin archivo, no por un dato declarado"""
    # Quién ya está dado de alta lo decide SIPP al capturar, así que aquí no se
    # puede afirmar que la carátula sea obligatoria: solo que podría hacer falta.
    lote = comun.lote()
    partidas = [comun.concepto("VIGILANCIA", 100.0)]
    transferencia = comun.solicitud(lote.id, "POR TRANSFERENCIA",
                                    partidas=partidas)
    efectivo = comun.solicitud(lote.id, "EN EFECTIVO",
                               partidas=[comun.concepto("VIGILANCIA", 200.0)],
                               forma_pago="Efectivo")
    assert documentos.falta_caratula(transferencia)
    assert not documentos.falta_caratula(efectivo), (
        "sin transferencia no hay cuenta que dar de alta")

    carpeta = comun.carpeta_temporal()
    documentos.registrar(transferencia.id,
                         comun.pdf_falso(carpeta, "c.pdf"),
                         documentos.TIPO_CARATULA, lote.id)
    assert not documentos.falta_caratula(transferencia)
