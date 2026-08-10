"""Ingesta: plantilla de Excel, lectura del archivo, CFDI y carátulas.

Lo que se prueba aquí es sobre todo **lo que el formato no debe perder**: la
CLABE con sus ceros, el importe con formato de moneda, el nombre con acentos.
Cada una de esas pérdidas manda dinero a otro lado o registra mal a una persona.
"""

from __future__ import annotations

import os

from openpyxl import Workbook, load_workbook

from core import db, plantilla_excel, validador
from core.adaptadores import caratulas
from core.adaptadores import cfdi as adaptador_cfdi
from core.adaptadores import excel as adaptador
from core.adaptadores import ocr_caratula
from core.db import CONCEPTO, INSUMO, Solicitud
from scripts.pruebas import comun

# Columnas de la hoja principal, en orden. Se arma a partir del contrato para
# que al agregar o quitar una columna las pruebas no queden desalineadas en
# silencio (pasó: se quitaron dos y todas las filas se corrieron dos lugares).
_CAMPOS = [c.clave for c in plantilla_excel.CAMPOS]


def _fila(**valores) -> list:
    """Fila de la hoja Solicitudes, por nombre de campo."""
    desconocidos = set(valores) - set(_CAMPOS)
    assert not desconocidos, f"campos que no existen: {desconocidos}"
    return [valores.get(c, "") for c in _CAMPOS]


def _plantilla_llena(filas: list[list], partidas: list[list] | None = None):
    """Genera la plantilla, le escribe filas y devuelve la ruta del archivo."""
    carpeta = comun.carpeta_temporal()
    ruta = os.path.join(carpeta, plantilla_excel.NOMBRE_ARCHIVO)
    plantilla_excel.generar(ruta)
    libro = load_workbook(ruta)
    hoja = libro[plantilla_excel.HOJA_SOLICITUDES]
    for i, fila in enumerate(filas, 3):        # la 2 es el ejemplo
        for j, valor in enumerate(fila, 1):
            hoja.cell(row=i, column=j).value = valor
    if partidas:
        hoja_p = libro[plantilla_excel.HOJA_PARTIDAS]
        for i, fila in enumerate(partidas, 3):
            for j, valor in enumerate(fila, 1):
                hoja_p.cell(row=i, column=j).value = valor
    lleno = os.path.join(carpeta, "lleno.xlsx")
    libro.save(lleno)
    return lleno


# --------------------------------------------------------------------------- #
#  Plantilla
# --------------------------------------------------------------------------- #
def probar_plantilla_trae_sus_hojas():
    """la plantilla trae instrucciones, solicitudes, partidas y catálogos"""
    ruta = os.path.join(comun.carpeta_temporal(),
                        plantilla_excel.NOMBRE_ARCHIVO)
    plantilla_excel.generar(ruta)
    libro = load_workbook(ruta)
    for hoja in (plantilla_excel.HOJA_INSTRUCCIONES,
                 plantilla_excel.HOJA_SOLICITUDES,
                 plantilla_excel.HOJA_PARTIDAS,
                 plantilla_excel.HOJA_CATALOGOS):
        assert hoja in libro.sheetnames, hoja
    hoja = libro[plantilla_excel.HOJA_SOLICITUDES]
    assert len(hoja.data_validations.dataValidation) >= 6, (
        "las columnas de catálogo deben traer desplegable")


def probar_plantilla_no_pregunta_por_el_alta_del_beneficiario():
    """la plantilla ya no pide «¿es nuevo?» ni su clave: lo resuelve el robot"""
    assert "beneficiario_nuevo" not in _CAMPOS
    assert "beneficiario_folio" not in _CAMPOS


def probar_la_fila_de_ejemplo_no_se_importa():
    """una plantilla recién bajada no da de alta nada"""
    ruta = os.path.join(comun.carpeta_temporal(),
                        plantilla_excel.NOMBRE_ARCHIVO)
    plantilla_excel.generar(ruta)
    importacion = adaptador.leer(ruta, "lote-x")
    assert not importacion.filas
    assert importacion.error


# --------------------------------------------------------------------------- #
#  Lectura
# --------------------------------------------------------------------------- #
def probar_lee_una_fila_completa():
    """una fila bien llena produce una solicitud válida"""
    lleno = _plantilla_llena([_fila(
        empresa="Abastecedora", sucursal="Corporativo",
        tipo_beneficiario="Acreedor", beneficiario_nombre="JUAN PEREZ LOPEZ",
        beneficiario_rfc="XAXX010101000", beneficiario_correo="j@x.invalid",
        forma_pago="Transferencia", tipo_gasto="No Deducible",
        cuenta_banco="BBVA", cuenta_clabe="012345678901234568",
        fecha_pago="15/09/2026", moneda="Pesos (MXN)",
        descripcion="Finiquito", concepto_nombre="PAGO PTU",
        importe="12500.00")])
    imp = adaptador.leer(lleno, "lote-x")
    assert len(imp.filas) == 1 and not imp.error, imp.error
    fila = imp.filas[0]
    assert fila.valida, fila.resumen_problemas
    assert fila.solicitud.importe_total == 12500.00
    assert fila.solicitud.origen == "EXCEL"


def probar_la_clabe_conserva_sus_ceros():
    """una CLABE que empieza en 0 no pierde dígitos"""
    # Si se convierte a número, «012...» queda en «12...» y el pago se va a otra
    # cuenta. Es la pérdida más cara del formato.
    lleno = _plantilla_llena([_fila(
        empresa="Abastecedora", sucursal="Corporativo",
        tipo_beneficiario="Acreedor", beneficiario_nombre="ANA LOPEZ",
        beneficiario_rfc="XAXX010101000", forma_pago="Transferencia",
        tipo_gasto="No Deducible", cuenta_clabe="012345678901234568",
        fecha_pago="15/09/2026", concepto_nombre="PAGO PTU", importe="100")])
    clabe = adaptador.leer(lleno, "l").filas[0].solicitud.cuenta_clabe
    assert clabe == "012345678901234568", clabe


def probar_importe_con_formato_de_moneda():
    """«$8,300.50» se lee como 8300.5"""
    lleno = _plantilla_llena([_fila(
        empresa="Abastecedora", sucursal="Corporativo",
        tipo_beneficiario="Acreedor", beneficiario_nombre="ANA LOPEZ",
        beneficiario_rfc="XAXX010101000", forma_pago="Transferencia",
        tipo_gasto="No Deducible", cuenta_clabe="012345678901234568",
        fecha_pago="15/09/2026", concepto_nombre="PAGO PTU",
        importe="$8,300.50")])
    assert adaptador.leer(lleno, "l").filas[0].solicitud.importe_total == 8300.50


def probar_la_clase_la_impone_el_tipo_de_beneficiario():
    """el archivo no elige la clase del renglón: la decide el tipo"""
    lleno = _plantilla_llena([
        _fila(empresa="Abastecedora", sucursal="Corporativo",
              tipo_beneficiario="Proveedor",
              beneficiario_nombre="PROVEEDOR SA",
              beneficiario_rfc="XAXX010101000", forma_pago="Transferencia",
              tipo_gasto="No Deducible", cuenta_clabe="012345678901234568",
              cuenta_banco="BBVA", fecha_pago="15/09/2026",
              concepto_nombre="Mantenimiento de equipo", importe="3000"),
        _fila(empresa="Abastecedora", sucursal="Corporativo",
              tipo_beneficiario="Acreedor", beneficiario_nombre="ACREEDOR SA",
              beneficiario_rfc="XAXX010101000", forma_pago="Transferencia",
              tipo_gasto="No Deducible", cuenta_clabe="012345678901234568",
              cuenta_banco="BBVA", fecha_pago="15/09/2026",
              concepto_nombre="PAGO PTU", importe="4000"),
    ])
    filas = adaptador.leer(lleno, "l").filas
    por_tipo = {f.solicitud.tipo_beneficiario: f for f in filas}
    prov = por_tipo["Proveedor"]
    assert prov.partidas[0].clase == INSUMO
    assert prov.partidas[0].insumo_nombre == "Mantenimiento de equipo"
    assert prov.solicitud.importe_total == 3000.0
    acre = por_tipo["Acreedor"]
    assert acre.partidas[0].clase == CONCEPTO
    assert acre.partidas[0].concepto_nombre == "PAGO PTU"


def probar_la_hoja_de_partidas_manda():
    """si se desglosa en la segunda hoja, el concepto de la fila se ignora"""
    lleno = _plantilla_llena(
        [_fila(referencia="SP-004", empresa="Abastecedora",
               sucursal="Corporativo", tipo_beneficiario="Acreedor",
               beneficiario_nombre="CONSTRUCTORA DEL NORTE",
               beneficiario_rfc="XAXX010101000", forma_pago="Transferencia",
               tipo_gasto="No Deducible", cuenta_clabe="012345678901234568",
               fecha_pago="15/09/2026", concepto_nombre="IGNORAR ESTE",
               importe="1")],
        partidas=[
            ["SP-004", "Concepto", "MANTENIMIENTO", "", "", "", "", "5000.00"],
            ["SP-004", "Concepto", "MATERIALES", "", "", "", "", "2500.00"],
            ["SP-004", "Insumo", "Pintura", "CC-10", "601-01", "1", "7500",
             "7500.00"],
        ])
    fila = adaptador.leer(lleno, "l").filas[0]
    conceptos = [p for p in fila.partidas if p.clase == CONCEPTO]
    assert len(conceptos) == 2
    assert all(p.concepto_nombre != "IGNORAR ESTE" for p in conceptos)
    assert fila.solicitud.importe_total == 7500.00


def probar_fila_incompleta_no_bloquea_a_las_demas():
    """se importa lo bueno y se marca lo malo"""
    lleno = _plantilla_llena([
        _fila(empresa="Abastecedora", sucursal="Corporativo",
              tipo_beneficiario="Acreedor", beneficiario_nombre="BUENA",
              beneficiario_rfc="XAXX010101000", forma_pago="Transferencia",
              tipo_gasto="No Deducible", cuenta_clabe="012345678901234568",
              cuenta_banco="BBVA", fecha_pago="15/09/2026",
              concepto_nombre="PAGO PTU", importe="1000"),
        # Sin empresa, CLABE corta y sin fecha.
        _fila(sucursal="Corporativo", tipo_beneficiario="Acreedor",
              beneficiario_nombre="MALA", beneficiario_rfc="XAXX010101000",
              forma_pago="Transferencia", tipo_gasto="No Deducible",
              cuenta_clabe="123", concepto_nombre="PAGO PTU", importe="100"),
    ])
    imp = adaptador.leer(lleno, "l")
    assert len(imp.validas) == 1 and len(imp.invalidas) == 1
    assert imp.invalidas[0].solicitud.beneficiario_nombre == "MALA"


def probar_duplicados_dentro_del_archivo():
    """dos filas que producen el mismo pago se detectan antes de guardar"""
    fila = _fila(empresa="Abastecedora", sucursal="Corporativo",
                 tipo_beneficiario="Acreedor", beneficiario_nombre="ANA",
                 beneficiario_rfc="XAXX010101000", forma_pago="Transferencia",
                 tipo_gasto="No Deducible", cuenta_clabe="012345678901234568",
                 fecha_pago="15/09/2026", concepto_nombre="PAGO PTU",
                 importe="100")
    imp = adaptador.leer(_plantilla_llena([fila, list(fila)]), "l")
    assert adaptador.detectar_duplicados(imp.filas)


def probar_encabezados_ajenos_piden_mapeo():
    """un archivo con otros nombres reconoce lo que puede y pide el resto"""
    carpeta = comun.carpeta_temporal()
    libro = Workbook()
    hoja = libro.active
    hoja.append(["EMPRESA", "PLAZA", "BENEFICIARIO", "MONTO", "COLUMNA RARA"])
    hoja.append(["Abastecedora", "Corporativo", "PEDRO SOSA", "500", "x"])
    ruta = os.path.join(carpeta, "ajeno.xlsx")
    libro.save(ruta)

    deteccion = adaptador.detectar(ruta)
    # Estos se reconocen por sinónimo; los demás hay que mapearlos a mano.
    assert "empresa" in deteccion.columnas
    assert "sucursal" in deteccion.columnas
    assert "beneficiario_nombre" in deteccion.columnas
    assert "importe" in deteccion.columnas
    assert "fecha_pago" in deteccion.faltantes


# --------------------------------------------------------------------------- #
#  CFDI
# --------------------------------------------------------------------------- #
_CFDI = """<?xml version="1.0" encoding="UTF-8"?>
<cfdi:Comprobante xmlns:cfdi="http://www.sat.gob.mx/cfd/{v}"
  xmlns:tfd="http://www.sat.gob.mx/TimbreFiscalDigital"
  Version="{ver}" Serie="A" Folio="1234" Fecha="2026-08-01T10:15:00"
  SubTotal="10000.00" Total="11600.00" Moneda="MXN" FormaPago="03">
  <cfdi:Emisor Rfc="CNO900101AAA" Nombre="CONSTRUCTORA DEL NORTE SA DE CV"/>
  <cfdi:Receptor Rfc="ASE010101BBB" Nombre="ADMINISTRACION DE SERVICIOS"/>
  <cfdi:Conceptos>
    <cfdi:Concepto ClaveProdServ="72101500" Cantidad="1" ClaveUnidad="E48"
      Descripcion="Mantenimiento preventivo" ValorUnitario="6000.00"
      Importe="6000.00"/>
    <cfdi:Concepto ClaveProdServ="72101501" Cantidad="2" ClaveUnidad="H87"
      Descripcion="Suministro de material" ValorUnitario="2000.00"
      Importe="4000.00"/>
  </cfdi:Conceptos>
  <cfdi:Impuestos TotalImpuestosTrasladados="1600.00"/>
  {complemento}
</cfdi:Comprobante>
"""
_TIMBRE = """<cfdi:Complemento>
    <tfd:TimbreFiscalDigital Version="1.1" UUID="{uuid}"
      FechaTimbrado="2026-08-01T10:16:00"/>
  </cfdi:Complemento>"""


def _xml(carpeta: str, nombre: str, *, version: str = "4",
         uuid: str = "A1B2C3D4-1111-2222-3333-444455556666",
         timbrado: bool = True) -> str:
    contenido = _CFDI.format(
        v=version, ver=f"{version}.0",
        complemento=_TIMBRE.format(uuid=uuid) if timbrado else "")
    ruta = os.path.join(carpeta, nombre)
    with open(ruta, "w", encoding="utf-8") as fh:
        fh.write(contenido)
    return ruta


def probar_cfdi_40():
    """se leen los datos de un CFDI 4.0 timbrado"""
    comp = adaptador_cfdi.leer_xml(_xml(comun.carpeta_temporal(), "f.xml"))
    assert comp.valido
    assert comp.uuid == "A1B2C3D4-1111-2222-3333-444455556666"
    assert comp.rfc_emisor == "CNO900101AAA"
    assert comp.fecha == "01/08/2026"
    assert comp.serie_folio == "A-1234"
    assert (comp.subtotal, comp.total, comp.impuestos) == (10000.0, 11600.0,
                                                           1600.0)
    assert comp.moneda == "Pesos (MXN)"
    assert comp.forma_pago == "Transferencia"
    assert len(comp.conceptos) == 2


def probar_cfdi_33():
    """la versión 3.3 se lee igual (cambia el espacio de nombres)"""
    comp = adaptador_cfdi.leer_xml(
        _xml(comun.carpeta_temporal(), "f33.xml", version="3"))
    assert comp.valido and len(comp.conceptos) == 2


def probar_cfdi_sin_timbre_se_rechaza():
    """sin UUID no sirve como respaldo de un pago"""
    comp = adaptador_cfdi.leer_xml(
        _xml(comun.carpeta_temporal(), "sin.xml", timbrado=False))
    assert not comp.valido and "UUID" in comp.error


def probar_xml_corrupto_no_tumba_la_carpeta():
    """un archivo ilegible se reporta y no interrumpe a los demás"""
    carpeta = comun.carpeta_temporal()
    with open(os.path.join(carpeta, "basura.xml"), "w", encoding="utf-8") as fh:
        fh.write("esto no es xml <<<")
    _xml(carpeta, "buena.xml")
    comprobantes = adaptador_cfdi.leer_carpeta(carpeta)
    assert len(comprobantes) == 2
    assert len([c for c in comprobantes if c.valido]) == 1


def probar_cfdi_a_solicitud():
    """los conceptos del CFDI se vuelven insumos, y el importe queda pendiente"""
    comp = adaptador_cfdi.leer_xml(_xml(comun.carpeta_temporal(), "f.xml"))
    solicitud, partidas = adaptador_cfdi.a_solicitud(
        comp, empresa="Abastecedora", sucursal="Corporativo")
    assert solicitud.beneficiario_rfc == "CNO900101AAA"
    assert solicitud.origen == "CFDI" and solicitud.tipo_gasto == "Deducible"
    # A propósito en cero: el concepto de pago no está en el CFDI.
    assert solicitud.importe_total == 0.0
    assert all(p.clase == INSUMO for p in partidas)
    assert not adaptador_cfdi.discrepancia_total(comp, partidas)
    partidas[0].importe = 5000.0
    assert adaptador_cfdi.discrepancia_total(comp, partidas)


def probar_pdf_se_empareja_por_uuid():
    """el UUID gana sobre el nombre del archivo"""
    carpeta = comun.carpeta_temporal()
    comp = adaptador_cfdi.leer_xml(_xml(carpeta, "factura.xml"))
    por_nombre = comun.pdf_falso(carpeta, "factura.pdf")
    por_uuid = comun.pdf_falso(carpeta, f"FACT_{comp.uuid}.pdf")
    assert adaptador_cfdi.emparejar_pdf(comp, [por_nombre, por_uuid]) == por_uuid
    assert adaptador_cfdi.emparejar_pdf(comp, [por_nombre]) == por_nombre


# --------------------------------------------------------------------------- #
#  Carátulas
# --------------------------------------------------------------------------- #
def probar_nombre_desde_el_archivo():
    """el nombre se extrae del archivo, descartando el ruido habitual"""
    casos = [
        ("CARATULA JUAN PEREZ LOPEZ BBVA.pdf", "JUAN PEREZ LOPEZ"),
        ("estado de cuenta - MARIA RUIZ SOTO - agosto 2026.pdf",
         "MARIA RUIZ SOTO"),
        ("Edo Cta Banorte LUIS ALBERTO DIAZ.jpg", "LUIS ALBERTO DIAZ"),
        ("caratula_ANA_LOPEZ_RUIZ_12345.pdf", "ANA LOPEZ RUIZ"),
        # Los acentos y la Ñ se CONSERVAN: el nombre se registra tal cual en
        # SIPP, y «JOSE NUNO» dejaría mal escrito a alguien para siempre.
        ("José Ñuño Gómez.png", "JOSÉ ÑUÑO GÓMEZ"),
        # Sin nada reconocible no se inventa un nombre.
        ("12345678.pdf", ""),
    ]
    for archivo, esperado in casos:
        obtenido = caratulas.nombre_desde_archivo(archivo)
        assert obtenido == esperado, f"{archivo}: «{obtenido}»"


def probar_una_caratula_es_una_solicitud():
    """cada archivo de la carpeta produce un borrador con su carátula"""
    carpeta = comun.carpeta_temporal()
    for nombre in ("CARATULA JUAN PEREZ LOPEZ BBVA.pdf",
                   "estado de cuenta - MARIA RUIZ SOTO - agosto.pdf",
                   "12345678.pdf"):
        comun.pdf_falso(carpeta, nombre)
    with open(os.path.join(carpeta, "no_es.txt"), "w") as fh:
        fh.write("x")

    lote = comun.lote()
    borradores = caratulas.crear_borradores(
        [carpeta], lote.id, empresa="Abastecedora", sucursal="Corporativo")
    assert len(borradores) == 3, "el .txt se ignora"
    # Nacen incompletos a propósito: les falta importe, fecha y concepto.
    assert all(not b.listo for b in borradores)


def probar_el_excel_completa_por_nombre():
    """el Excel rellena lo que falta, emparejando aunque cambie el orden"""
    carpeta = comun.carpeta_temporal()
    comun.pdf_falso(carpeta, "CARATULA JUAN PEREZ LOPEZ.pdf")
    comun.pdf_falso(carpeta, "CARATULA MARIA RUIZ SOTO.pdf")
    lote = comun.lote()
    borradores = caratulas.crear_borradores(
        [carpeta], lote.id, empresa="Abastecedora", sucursal="Corporativo")

    lleno = _plantilla_llena([
        # Nombre INVERTIDO: debe emparejar igual.
        _fila(empresa="Abastecedora", sucursal="Corporativo",
              tipo_beneficiario="Acreedor",
              beneficiario_nombre="RUIZ SOTO MARIA",
              beneficiario_rfc="XAXX010101000", forma_pago="Transferencia",
              tipo_gasto="No Deducible", cuenta_clabe="002345678901234565",
              fecha_pago="15/09/2026", concepto_nombre="VIGILANCIA",
              importe="8300.50"),
        _fila(empresa="Abastecedora", sucursal="Corporativo",
              tipo_beneficiario="Acreedor",
              beneficiario_nombre="JUAN PEREZ LOPEZ",
              beneficiario_rfc="XAXX010101000", forma_pago="Transferencia",
              tipo_gasto="No Deducible", cuenta_clabe="012345678901234568",
              fecha_pago="15/09/2026", concepto_nombre="VIGILANCIA",
              importe="12500.00"),
        _fila(empresa="Abastecedora", sucursal="Corporativo",
              tipo_beneficiario="Acreedor",
              beneficiario_nombre="SIN CARATULA PEREZ",
              beneficiario_rfc="XAXX010101000", forma_pago="Transferencia",
              tipo_gasto="No Deducible", cuenta_clabe="012345678901234597",
              fecha_pago="15/09/2026", concepto_nombre="VIGILANCIA",
              importe="999.00"),
    ])
    filas = adaptador.leer(lleno, lote.id).filas
    resumen = caratulas.completar_con_excel(borradores, filas)

    assert resumen["emparejados"] == 2, resumen
    assert "SIN CARATULA PEREZ" in resumen["sin_caratula"]
    maria = [b for b in borradores if "MARIA" in b.nombre_detectado][0]
    # El nombre de la CARÁTULA prevalece: es la que acredita la cuenta.
    assert maria.nombre_detectado == "MARIA RUIZ SOTO"
    assert maria.solicitud.importe_total == 8300.50
    juan = [b for b in borradores
            if b.nombre_detectado == "JUAN PEREZ LOPEZ"][0]
    assert juan.listo and juan.solicitud.cuenta_clabe == "012345678901234568"


def probar_una_fila_del_excel_se_usa_una_vez():
    """dos carátulas no pueden alimentarse del mismo renglón"""
    carpeta = comun.carpeta_temporal()
    comun.pdf_falso(carpeta, "CARATULA JUAN PEREZ LOPEZ.pdf")
    comun.pdf_falso(carpeta, "CARATULA JUAN PEREZ LOPEZ copia.pdf")
    lote = comun.lote()
    borradores = caratulas.crear_borradores([carpeta], lote.id)
    lleno = _plantilla_llena([_fila(
        empresa="Abastecedora", sucursal="Corporativo",
        tipo_beneficiario="Acreedor", beneficiario_nombre="JUAN PEREZ LOPEZ",
        beneficiario_rfc="XAXX010101000", forma_pago="Transferencia",
        tipo_gasto="No Deducible", cuenta_clabe="012345678901234568",
        fecha_pago="15/09/2026", concepto_nombre="VIGILANCIA",
        importe="100")])
    resumen = caratulas.completar_con_excel(
        borradores, adaptador.leer(lleno, lote.id).filas)
    assert resumen["emparejados"] == 1, resumen


# --------------------------------------------------------------------------- #
#  Lectura de la carátula (OCR)
# --------------------------------------------------------------------------- #
# La CLABE es el dato del que depende a quién se le paga, y aquí la escribe un
# OCR, no una persona. Por eso se prueba el verificador antes que nada: es lo
# único que distingue una CLABE bien leída de una con un dígito confundido, que
# tiene exactamente la misma forma.
def probar_la_clabe_se_valida_con_su_digito_verificador():
    """una CLABE con un dígito cambiado se detecta, aunque mida 18"""
    buena = "012345678901234568"
    assert ocr_caratula.clabe_valida(buena)
    # Mismo largo y mismo prefijo de banco, el ÚLTIMO dígito cambiado: se ve
    # igual de bien que la buena y la única forma de cazarla es el verificador.
    assert not ocr_caratula.clabe_valida("012345678901234567")
    # Y un dígito cambiado a media CLABE, que es el error típico del OCR.
    assert not ocr_caratula.clabe_valida("012345678901294568")
    assert not ocr_caratula.clabe_valida("01234567890123456")   # 17 dígitos
    assert not ocr_caratula.clabe_valida("")
    # El adaptador solo reexporta la del validador: tiene que ser la MISMA
    # función, o con el tiempo una aceptaría lo que la otra rechaza.
    assert ocr_caratula.clabe_valida is validador.clabe_valida


def probar_se_lee_la_clabe_el_titular_y_el_banco():
    """de una carátula se sacan los datos que la solicitud necesita"""
    texto = (
        "BBVA MEXICO S.A.\n"
        "Estado de cuenta\n"
        "Titular de la cuenta: MARIA DE LA CRUZ ÑUÑO\n"
        "No. de cuenta: 1234567890\n"
        "CLABE interbancaria: 012345678901234568\n")
    datos = ocr_caratula.interpretar(texto)
    assert datos.clabe == "012345678901234568"
    assert datos.clabe_confiable
    # El banco sale del prefijo de la CLABE, más confiable que leer el logo.
    assert datos.banco == "BBVA"
    # Los acentos y la Ñ se conservan: así se registra en SIPP para siempre.
    assert datos.titular == "MARIA DE LA CRUZ ÑUÑO", datos.titular
    assert datos.cuenta == "1234567890"


def probar_no_se_devuelve_el_rfc_del_banco():
    """el RFC del encabezado es de la institución, no del beneficiario"""
    # Aparece en casi toda carátula y es el falso positivo más frecuente:
    # tomarlo daría de alta al beneficiario con el RFC del banco.
    texto = (
        "BANCO SANTANDER MEXICO, INSTITUCION DE BANCA MULTIPLE\n"
        "R.F.C. BSM970519DU8\n"
        "Titular: JUAN PEREZ LOPEZ\n"
        "CLABE: 014111222333444559\n")
    datos = ocr_caratula.interpretar(texto)
    assert datos.rfc == "", f"se coló el RFC del banco: {datos.rfc}"
    assert datos.titular == "JUAN PEREZ LOPEZ"


def probar_el_titular_de_la_caratula_manda_sobre_el_nombre_del_archivo():
    """el beneficiario sale de la carátula, no de cómo se llame el archivo"""
    # Es el cambio de fondo de esta función. El nombre del archivo lo escribió
    # alguien de memoria al armar la carpeta; el titular es lo que el banco
    # acredita como dueño de la cuenta a la que se va a pagar.
    carpeta = comun.carpeta_temporal()
    comun.caratula_pdf(carpeta, "CARATULA JUAN PEREZ LOPEZ.pdf",
                       titular="MARIA RUIZ SOTO", clabe="012345678901234568")
    lote = comun.lote()
    borradores = caratulas.crear_borradores([carpeta], lote.id)

    assert len(borradores) == 1
    b = borradores[0]
    assert b.solicitud.beneficiario_nombre == "MARIA RUIZ SOTO", \
        b.solicitud.beneficiario_nombre
    assert b.origen_nombre == caratulas.NOMBRE_DE_OCR
    # El nombre de la cuenta ES el del beneficiario.
    assert b.solicitud.cuenta_titular == "MARIA RUIZ SOTO"
    # Y de paso llegan la CLABE y el banco, que el archivo no podía dar.
    assert b.solicitud.cuenta_clabe == "012345678901234568"
    assert b.solicitud.cuenta_banco == "BBVA"
    assert b.clabe_confiable


def probar_una_caratula_ilegible_no_tumba_la_carga():
    """un archivo que no se puede leer cae al nombre del archivo, no falla"""
    carpeta = comun.carpeta_temporal()
    comun.pdf_falso(carpeta, "CARATULA JUAN PEREZ LOPEZ.pdf")
    lote = comun.lote()
    borradores = caratulas.crear_borradores([carpeta], lote.id)
    assert len(borradores) == 1
    b = borradores[0]
    assert b.solicitud.beneficiario_nombre == "JUAN PEREZ LOPEZ"
    assert b.origen_nombre == caratulas.NOMBRE_DE_ARCHIVO
    # El titular de la cuenta acompaña al beneficiario: son la misma persona.
    assert b.solicitud.cuenta_titular == "JUAN PEREZ LOPEZ"


# --------------------------------------------------------------------------- #
#  Emparejamiento con el Excel por CLABE
# --------------------------------------------------------------------------- #
def _borrador(nombre: str, clabe: str) -> caratulas.Borrador:
    """Borrador ya «leído», para probar el emparejamiento sin OCR ni archivos."""
    return caratulas.Borrador(
        ruta=f"{nombre}.pdf", nombre_detectado=nombre,
        origen_nombre=caratulas.NOMBRE_DE_OCR,
        solicitud=Solicitud(
            empresa="Abastecedora", sucursal="Corporativo",
            tipo_beneficiario="Acreedor", beneficiario_nombre=nombre,
            cuenta_titular=nombre, cuenta_clabe=clabe,
            forma_pago="Transferencia", tipo_gasto="No Deducible"))


def probar_el_excel_empareja_por_clabe_aunque_el_nombre_difiera():
    """la CLABE manda: es el mismo número en los dos lados"""
    # El caso real: en el Excel el nombre viene invertido, abreviado o con un
    # apellido de más. Emparejar por nombre lo dejaría fuera; por CLABE, no.
    lote = comun.lote()
    borradores = [_borrador("MARIA RUIZ SOTO", "012345678901234568")]
    lleno = _plantilla_llena([_fila(
        empresa="Abastecedora", sucursal="Corporativo",
        tipo_beneficiario="Acreedor",
        beneficiario_nombre="MA. RUIZ S. DE JESUS",
        beneficiario_rfc="XAXX010101000", forma_pago="Transferencia",
        tipo_gasto="No Deducible", cuenta_clabe="012345678901234568",
        fecha_pago="15/09/2026", concepto_nombre="VIGILANCIA",
        importe="8300.50")])
    resumen = caratulas.completar_con_excel(
        borradores, adaptador.leer(lleno, lote.id).filas)

    assert resumen["por_clabe"] == 1, resumen
    assert resumen["por_nombre"] == 0, resumen
    b = borradores[0]
    assert b.solicitud.importe_total == 8300.50
    # El nombre NO se sobrescribe: manda el de la carátula, que es la que
    # acredita la cuenta.
    assert b.solicitud.beneficiario_nombre == "MARIA RUIZ SOTO"


def probar_sin_clabe_se_empareja_por_nombre():
    """cuando el OCR no pudo leer la CLABE, el nombre sigue sirviendo"""
    lote = comun.lote()
    borradores = [_borrador("JUAN PEREZ LOPEZ", "")]
    lleno = _plantilla_llena([_fila(
        empresa="Abastecedora", sucursal="Corporativo",
        tipo_beneficiario="Acreedor", beneficiario_nombre="PEREZ LOPEZ JUAN",
        beneficiario_rfc="XAXX010101000", forma_pago="Transferencia",
        tipo_gasto="No Deducible", cuenta_clabe="072987654321098764",
        fecha_pago="15/09/2026", concepto_nombre="VIGILANCIA",
        importe="100")])
    resumen = caratulas.completar_con_excel(
        borradores, adaptador.leer(lleno, lote.id).filas)

    assert resumen["por_nombre"] == 1, resumen
    # Al no tenerla, la CLABE del Excel sí completa el hueco.
    assert borradores[0].solicitud.cuenta_clabe == "072987654321098764"


def probar_clabe_distinta_conserva_la_de_la_caratula_y_avisa():
    """si las dos CLABE existen y difieren, manda la carátula y queda constancia"""
    # Casi siempre significa que el Excel trae el renglón de otra persona. No se
    # bloquea —el usuario decide— pero no puede pasar en silencio.
    lote = comun.lote()
    borradores = [_borrador("ANA LOPEZ RUIZ", "012345678901234568")]
    lleno = _plantilla_llena([_fila(
        empresa="Abastecedora", sucursal="Corporativo",
        tipo_beneficiario="Acreedor", beneficiario_nombre="ANA LOPEZ RUIZ",
        beneficiario_rfc="XAXX010101000", forma_pago="Transferencia",
        tipo_gasto="No Deducible", cuenta_clabe="072987654321098764",
        fecha_pago="15/09/2026", concepto_nombre="VIGILANCIA",
        importe="500")])
    resumen = caratulas.completar_con_excel(
        borradores, adaptador.leer(lleno, lote.id).filas)

    assert resumen["discrepancias"] == 1, resumen
    b = borradores[0]
    assert b.solicitud.cuenta_clabe == "012345678901234568", "manda la carátula"
    assert b.avisos and "no es la de la carátula" in b.avisos[0]
    # Se empareja igual: el resto de los datos del renglón sí sirven.
    assert b.solicitud.importe_total == 500.0


def probar_la_clabe_gana_la_fila_antes_que_el_nombre():
    """una coincidencia por nombre no puede robarse la fila que casa por CLABE"""
    # Por eso el emparejamiento va en dos pasadas COMPLETAS y no solicitud por
    # solicitud: resolviendo una a una, la homónima sin CLABE se llevaría el
    # único renglón y la que sí casaba por CLABE se quedaría sin él.
    lote = comun.lote()
    con_clabe = _borrador("JUAN PEREZ LOPEZ", "012345678901234568")
    sin_clabe = _borrador("JUAN PEREZ LOPEZ", "")
    lleno = _plantilla_llena([_fila(
        empresa="Abastecedora", sucursal="Corporativo",
        tipo_beneficiario="Acreedor", beneficiario_nombre="JUAN PEREZ LOPEZ",
        beneficiario_rfc="XAXX010101000", forma_pago="Transferencia",
        tipo_gasto="No Deducible", cuenta_clabe="012345678901234568",
        fecha_pago="15/09/2026", concepto_nombre="VIGILANCIA",
        importe="750")])
    # El que NO tiene CLABE va primero en la lista, para que ganar por orden no
    # baste: tiene que ganar por criterio.
    caratulas.completar_con_excel([sin_clabe, con_clabe],
                                  adaptador.leer(lleno, lote.id).filas)

    assert con_clabe.emparejado_por == caratulas.POR_CLABE
    assert con_clabe.solicitud.importe_total == 750.0
    assert sin_clabe.emparejado_por == "", "no debía quedarse con esa fila"
