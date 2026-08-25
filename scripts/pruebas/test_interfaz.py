"""Lógica de las pantallas y los modales, contra una página simulada.

Estas pruebas comprueban el COMPORTAMIENTO: que la vista previa no se pueda
saltar, que el editor cambie según el tipo de beneficiario, que un guardado
incompleto se bloquee. Que además se DIBUJEN es otra cosa, y la cubre
`scripts/smoke_render.py`, que abre la ventana de verdad.
"""

from __future__ import annotations

import os

from openpyxl import load_workbook

from core import asignacion, conceptos, db, documentos, plantilla_excel
from core.db import CONCEPTO, INSUMO
from scripts.pruebas import comun


# --------------------------------------------------------------------------- #
#  Pantallas
# --------------------------------------------------------------------------- #
def probar_las_pantallas_se_construyen():
    """cada pantalla se arma y expone su contenido sin abrir ventana"""
    from ui.bitacora import SeccionBitacora
    from ui.conceptos import SeccionConceptos
    from ui.documentos import SeccionDocumentos
    from ui.solicitudes import SeccionSolicitudes

    app = comun.AppFalsa()
    for clase in (SeccionDocumentos, SeccionSolicitudes, SeccionConceptos,
                  SeccionBitacora):
        pantalla = clase(app)
        assert pantalla.contenido is not None, clase.__name__


def probar_configuracion_dice_contra_que_sipp_se_trabaja():
    """el modal enseña la URL real y advierte que cada folio es de verdad"""
    # Ya no se elige ambiente, pero SÍ hay que poder comprobar contra dónde se
    # está trabajando: es lo que separa un ensayo de una captura en el ERP real.
    from core.catalogos import AMBIENTES
    from ui.configuracion import SeccionConfiguracion

    app = comun.AppFalsa()
    modal = SeccionConfiguracion(app)
    assert modal.dialogo is not None
    assert modal.txt_url.value == AMBIENTES["PRODUCCION"]
    assert modal.txt_aviso_prod.visible, "el aviso del folio real siempre a la vista"
    assert not hasattr(modal, "dd_ambiente"), "no debe poder elegirse"


def probar_el_ambiente_no_se_puede_dejar_apuntando_a_pruebas():
    """una preferencia vieja no puede mandar el lote al SIPP equivocado"""
    # Se podía elegir, y la elección quedaba guardada en el equipo: haberla
    # dejado en pruebas una vez bastaba para capturar un lote entero contra
    # stage sin que nada lo delatara.
    from core import preferencias
    from ui.configuracion import ambiente_actual

    preferencias.guardar_valor("ambiente", "PRUEBAS")
    try:
        assert ambiente_actual() == "PRODUCCION"
    finally:
        preferencias.guardar_valor("ambiente", "PRODUCCION")


def probar_solicitudes_crea_lote_y_da_de_alta():
    """la pantalla crea un lote si no hay y guarda lo que le pasa el modal"""
    from ui.solicitudes import SeccionSolicitudes

    app = comun.AppFalsa()
    pantalla = SeccionSolicitudes(app)
    pantalla.cargar_desde_db()
    assert pantalla._lote is not None
    assert "0 solicitudes" in pantalla.txt_resumen.value

    lote_id = pantalla._lote.id
    solicitud = db.Solicitud(
        lote_id=lote_id, empresa="Abastecedora", sucursal="Corporativo",
        tipo_beneficiario="Acreedor", beneficiario_nombre="ANA LOPEZ",
        beneficiario_rfc="XAXX010101000", cuenta_clabe="012345678901234568",
        cuenta_banco="BBVA", forma_pago="Transferencia",
        tipo_gasto="No Deducible", fecha_pago=comun.fecha_futura())
    pantalla._recibir_solicitud(solicitud, [comun.concepto("VIGILANCIA", 1200.0)])
    assert len(pantalla._solicitudes) == 1
    assert "1 solicitud" in pantalla.txt_resumen.value


def probar_desplegar_el_detalle_de_una_fila():
    """la fila maestro-detalle se abre sin error"""
    from ui.solicitudes import SeccionSolicitudes

    app = comun.AppFalsa()
    pantalla = SeccionSolicitudes(app)
    pantalla.cargar_desde_db()
    s = comun.solicitud(pantalla._lote.id, "ANA LOPEZ",
                        partidas=[comun.concepto("VIGILANCIA", 100.0),
                                  comun.insumo("Servicio", 100.0)])
    pantalla._recargar_solicitudes()
    pantalla._alternar_detalle(s.id)
    assert s.id in pantalla._expandidos
    pantalla._alternar_detalle(s.id)
    assert s.id not in pantalla._expandidos


def probar_duplicar_precarga_el_editor():
    """duplicar no guarda: abre el editor con una copia"""
    # Una copia idéntica produce la MISMA clave de idempotencia y sería
    # rechazada; duplicar sirve para partir de una parecida y cambiarle algo.
    from ui.solicitudes import SeccionSolicitudes

    app = comun.AppFalsa()
    pantalla = SeccionSolicitudes(app)
    pantalla.cargar_desde_db()
    s = comun.solicitud(pantalla._lote.id, "ANA LOPEZ",
                        partidas=[comun.concepto("VIGILANCIA", 100.0)])
    pantalla._recargar_solicitudes()
    # Se duplica por su id, desde el ícono de la fila: ya no hace falta
    # seleccionarla antes.
    pantalla._duplicar(s.id)
    assert pantalla.captura.tf_nombre.value == "ANA LOPEZ"
    assert pantalla.captura._solicitud.id != s.id
    assert len(db.listar_solicitudes(pantalla._lote.id)) == 1, "no debe guardar"


def probar_deshacer_por_fila_no_toca_al_resto():
    """deshacer en una solicitud la revierte solo a ella, y deja deshacer más"""
    from ui.solicitudes import SeccionSolicitudes

    app = comun.AppFalsa()
    pantalla = SeccionSolicitudes(app)
    pantalla.cargar_desde_db()
    lote_id = pantalla._lote.id
    ana = comun.solicitud(lote_id, "ANA LOPEZ", partidas=[
        comun.insumo("Facturado", 1000.0, origen="CFDI")])
    luis = comun.solicitud(lote_id, "LUIS DIAZ MORA", partidas=[
        comun.insumo("Facturado", 500.0, origen="CFDI")])
    pantalla._recargar_solicitudes()
    # Sin operación masiva no hay estado anterior al que volver.
    assert not db.solicitud_en_snapshot(lote_id, ana.id)

    plan = asignacion.calcular(lote_id, nombre="VIGILANCIA")
    asignacion.aplicar(lote_id, plan)
    pantalla._recargar_solicitudes()
    assert db.solicitud_en_snapshot(lote_id, ana.id)
    assert len(db.listar_partidas(ana.id)) == 2, "la masiva agregó un renglón"

    # Revertir a ANA no revierte a LUIS.
    pantalla._deshacer_fila(ana.id)
    comun.confirmar_dialogo(app.page)
    assert len(db.listar_partidas(ana.id)) == 1
    assert len(db.listar_partidas(luis.id)) == 2, "el resto del lote no se toca"

    # Y el snapshot SIGUE ahí: si consumirlo fuera lo normal, la segunda
    # solicitud se quedaría sin vuelta atrás por haber arreglado la primera.
    assert db.solicitud_en_snapshot(lote_id, luis.id)
    pantalla._deshacer_fila(luis.id)
    comun.confirmar_dialogo(app.page)
    assert len(db.listar_partidas(luis.id)) == 1


def probar_el_check_general_selecciona_todo_el_lote():
    """el check del encabezado marca y desmarca todas las solicitudes"""
    from ui.solicitudes import SeccionSolicitudes

    app = comun.AppFalsa()
    pantalla = SeccionSolicitudes(app)
    pantalla.cargar_desde_db()
    for nombre in ("ANA LOPEZ", "LUIS DIAZ", "MARIA RUIZ"):
        comun.solicitud(pantalla._lote.id, nombre,
                        partidas=[comun.concepto("VIGILANCIA", 100.0)])
    pantalla._recargar_solicitudes()

    assert not pantalla._seleccionados
    assert pantalla.barra_global.visible and not pantalla.barra_seleccion.visible

    pantalla._alternar_todos(comun.evento_check(True))
    assert len(pantalla._seleccionados) == 3
    # Las dos barras se excluyen: con selección se ve la contextual y no la
    # global, que es de lo que depende que no queden botones muertos a la vista.
    assert pantalla.barra_seleccion.visible
    assert not pantalla.barra_global.visible
    assert pantalla.txt_seleccion.value == "3 seleccionadas"

    pantalla._alternar_todos(comun.evento_check(False))
    assert not pantalla._seleccionados
    assert pantalla.barra_global.visible


def probar_el_check_general_refleja_la_seleccion_manual():
    """marcar las filas una por una acaba marcando el check del encabezado"""
    # Si no, quedaría vacío con todo el lote seleccionado, diciendo lo contrario
    # de lo que pasa.
    from ui.solicitudes import SeccionSolicitudes

    app = comun.AppFalsa()
    pantalla = SeccionSolicitudes(app)
    pantalla.cargar_desde_db()
    ids = [comun.solicitud(pantalla._lote.id, n,
                           partidas=[comun.concepto("VIGILANCIA", 100.0)]).id
           for n in ("ANA LOPEZ", "LUIS DIAZ")]
    pantalla._recargar_solicitudes()

    pantalla._alternar_seleccion(ids[0], comun.evento_check(True))
    assert not pantalla.chk_todos.value, "falta una por marcar"
    pantalla._alternar_seleccion(ids[1], comun.evento_check(True))
    assert pantalla.chk_todos.value, "ya están todas marcadas"


def probar_eliminar_desde_la_barra_borra_lo_seleccionado():
    """el botón de la barra borra la selección, no una solicitud suelta"""
    # Esta prueba existe por un error real: el manejador recibía el EVENTO de
    # Flet en el parámetro del id y trataba de borrar una solicitud con ese
    # objeto como identificador.
    from ui.solicitudes import SeccionSolicitudes

    app = comun.AppFalsa()
    pantalla = SeccionSolicitudes(app)
    pantalla.cargar_desde_db()
    lote_id = pantalla._lote.id
    ana = comun.solicitud(lote_id, "ANA LOPEZ",
                          partidas=[comun.concepto("VIGILANCIA", 100.0)])
    luis = comun.solicitud(lote_id, "LUIS DIAZ",
                           partidas=[comun.concepto("VIGILANCIA", 100.0)])
    pantalla._recargar_solicitudes()
    pantalla._seleccionados = {ana.id}

    # Se pulsa el BOTÓN, no su manejador, y se le pasa el evento como hace
    # Flet: el error consistía precisamente en que ese evento aterrizaba en el
    # parámetro del id, así que llamar al método directo no lo habría visto.
    boton = comun.boton_por_texto(pantalla.barra_seleccion, "Eliminar")
    boton.on_click(comun.evento_check(True))
    comun.confirmar_dialogo(app.page)

    quedan = [s.id for s in db.listar_solicitudes(lote_id)]
    assert quedan == [luis.id], quedan
    assert not pantalla._seleccionados


def probar_eliminar_una_fila_no_necesita_seleccionarla():
    """el ícono de la fila borra esa solicitud sin tocar la selección"""
    from ui.solicitudes import SeccionSolicitudes

    app = comun.AppFalsa()
    pantalla = SeccionSolicitudes(app)
    pantalla.cargar_desde_db()
    lote_id = pantalla._lote.id
    ana = comun.solicitud(lote_id, "ANA LOPEZ",
                          partidas=[comun.concepto("VIGILANCIA", 100.0)])
    luis = comun.solicitud(lote_id, "LUIS DIAZ",
                           partidas=[comun.concepto("VIGILANCIA", 100.0)])
    pantalla._recargar_solicitudes()
    # Con OTRA solicitud seleccionada: borrar por fila no debe confundirse con
    # la selección viva, que es justo lo que pasaba al tener que marcarla antes.
    pantalla._seleccionados = {luis.id}

    pantalla._eliminar(ana.id)
    comun.confirmar_dialogo(app.page)

    quedan = [s.id for s in db.listar_solicitudes(lote_id)]
    assert quedan == [luis.id], quedan
    assert pantalla._seleccionados == {luis.id}, "la selección no se toca"


def probar_nueva_solicitud_pregunta_la_fuente():
    """«Nueva solicitud» ofrece las fuentes y abre la que se elija"""
    from ui.solicitudes import SeccionSolicitudes

    app = comun.AppFalsa()
    pantalla = SeccionSolicitudes(app)
    pantalla.cargar_desde_db()

    pantalla._nueva_solicitud()
    assert app.page.dialogos, "debe preguntar de dónde sale la solicitud"

    # La segunda opción es el alta desde carátulas; elegirla cierra el diálogo
    # de fuentes y abre ese modal.
    comun.elegir_del_dialogo(app.page, 1)
    assert app.page.dialogos[-1] is pantalla.alta_caratulas.modal.dialogo


# --------------------------------------------------------------------------- #
#  Modal de captura
# --------------------------------------------------------------------------- #
def probar_captura_bloquea_lo_incompleto():
    """guardar con datos faltantes muestra los hallazgos y no cierra"""
    from ui.captura_solicitud import CapturaSolicitud

    app = comun.AppFalsa()
    guardadas = []
    captura = CapturaSolicitud(app, lambda s, p: guardadas.append(s))
    lote = comun.lote()
    captura.abrir(lote.id)
    captura._guardar()
    assert captura.txt_hallazgos.visible
    assert not guardadas


def probar_captura_calcula_el_total():
    """el total se recalcula al editar un renglón"""
    from ui.captura_solicitud import CapturaSolicitud

    app = comun.AppFalsa()
    captura = CapturaSolicitud(app, lambda *_a: None)
    captura.abrir(comun.lote().id)
    renglon = captura.ed_conceptos.renglones[0]
    renglon.tf_nombre.value = "VIGILANCIA"
    renglon.tf_importe.value = "2,500.75"
    captura.ed_conceptos.recalcular()
    assert captura.txt_total.value == "$2,500.75", captura.txt_total.value
    partida = renglon.recolectar()
    assert partida.concepto_nombre == "VIGILANCIA" and partida.importe == 2500.75


def probar_el_editor_cambia_con_el_tipo_de_beneficiario():
    """a un Proveedor se le muestran insumos; a un Acreedor, conceptos"""
    from ui.captura_solicitud import CapturaSolicitud

    app = comun.AppFalsa()
    captura = CapturaSolicitud(app, lambda *_a: None)
    captura.abrir(comun.lote().id)
    assert captura.ed_conceptos.control.visible
    assert not captura.ed_insumos.control.visible

    captura.dd_tipo_ben.value = "Proveedor"
    captura._cambio_tipo_beneficiario()
    assert captura.ed_insumos.control.visible
    assert not captura.ed_conceptos.control.visible
    # Y solo se recolecta el desglose que aplica.
    _s, partidas = captura._recolectar()
    assert all(p.clase == INSUMO for p in partidas)


def probar_captura_ofrece_el_catalogo_de_conceptos():
    """el renglón de concepto es un desplegable editable con el catálogo"""
    from ui.captura_solicitud import CapturaSolicitud

    conceptos.importar(["VIGILANCIA", "NO DEDUCIBLE"], "Abastecedora")
    app = comun.AppFalsa()
    captura = CapturaSolicitud(app, lambda *_a: None)
    captura.abrir(comun.lote().id)
    renglon = captura.ed_conceptos.renglones[0]
    opciones = [o.key for o in renglon.tf_nombre.options]
    assert "VIGILANCIA" in opciones
    assert renglon.tf_nombre.editable, "debe dejar escribir uno que no esté"


def probar_captura_muestra_los_adjuntos_al_editar():
    """al reabrir una solicitud se ven sus archivos"""
    # Esta es la rama donde vivía un NameError: solo se ejecuta cuando la
    # solicitud YA trae archivos, y abrir el formulario en alta nunca la tocaba.
    from ui.captura_solicitud import CapturaSolicitud

    lote = comun.lote()
    s = comun.solicitud(lote.id, "ANA LOPEZ",
                        partidas=[comun.concepto("VIGILANCIA", 100.0)])
    carpeta = comun.carpeta_temporal()
    documentos.registrar(s.id, comun.pdf_falso(carpeta, "caratula.pdf"),
                         documentos.TIPO_CARATULA, lote.id)
    documentos.registrar(s.id, comun.pdf_falso(carpeta, "vobo.pdf"),
                         documentos.TIPO_VOBO, lote.id)

    app = comun.AppFalsa()
    captura = CapturaSolicitud(app, lambda *_a: None)
    captura.abrir(lote.id, s, db.listar_partidas(s.id))
    assert "caratula.pdf" in (captura.txt_caratula.value or "")
    assert "vobo.pdf" in (captura.txt_vobo.value or "")


# --------------------------------------------------------------------------- #
#  Modal de carga masiva
# --------------------------------------------------------------------------- #
def _excel_de_prueba(carpeta: str) -> str:
    campos = [c.clave for c in plantilla_excel.CAMPOS]

    def fila(**valores):
        return [valores.get(c, "") for c in campos]

    ruta = os.path.join(carpeta, plantilla_excel.NOMBRE_ARCHIVO)
    plantilla_excel.generar(ruta)
    libro = load_workbook(ruta)
    hoja = libro[plantilla_excel.HOJA_SOLICITUDES]
    filas = [
        fila(empresa="Abastecedora", sucursal="Corporativo",
             tipo_beneficiario="Acreedor", beneficiario_nombre="ANA LOPEZ RUIZ",
             beneficiario_rfc="XAXX010101000", forma_pago="Transferencia",
             tipo_gasto="No Deducible", cuenta_banco="BBVA",
             cuenta_clabe="012345678901234568", fecha_pago=comun.fecha_futura(),
             concepto_nombre="VIGILANCIA", importe="9500"),
        fila(empresa="Abastecedora", sucursal="Corporativo",
             tipo_beneficiario="Acreedor", beneficiario_nombre="LUIS DIAZ MORA",
             beneficiario_rfc="XAXX010101000", forma_pago="Transferencia",
             tipo_gasto="No Deducible", cuenta_banco="BBVA",
             cuenta_clabe="012345678901234568", fecha_pago=comun.fecha_futura(),
             concepto_nombre="VIGILANCIA", importe="7200"),
        # Sin empresa y con CLABE corta: debe salir marcada.
        fila(sucursal="Corporativo", tipo_beneficiario="Acreedor",
             beneficiario_nombre="SIN EMPRESA",
             beneficiario_rfc="XAXX010101000", forma_pago="Transferencia",
             tipo_gasto="No Deducible", cuenta_clabe="12",
             concepto_nombre="VIGILANCIA", importe="100"),
    ]
    for i, f in enumerate(filas, 3):
        for j, v in enumerate(f, 1):
            hoja.cell(row=i, column=j).value = v
    lleno = os.path.join(carpeta, "lleno.xlsx")
    libro.save(lleno)
    return lleno


def probar_carga_masiva_exige_revisar_antes():
    """la previa está oculta y el botón bloqueado hasta analizar el archivo"""
    from ui.carga_masiva import CargaMasiva

    app = comun.AppFalsa()
    carga = CargaMasiva(app, lambda *_a: None)
    carga.abrir(comun.lote().id)
    assert not carga.previa.visible
    assert carga.btn_importar.disabled


def probar_carga_masiva_importa_lo_bueno():
    """se importan las filas correctas y se marcan las malas"""
    from ui.carga_masiva import CargaMasiva

    app = comun.AppFalsa()
    importadas = []
    carga = CargaMasiva(app, lambda n: importadas.append(n))
    lote = comun.lote()
    carga.abrir(lote.id)
    carga._ruta = _excel_de_prueba(comun.carpeta_temporal())
    comun.correr(carga._analizar())

    assert carga.previa.visible
    assert carga.btn_importar.text == "Importar 2", carga.btn_importar.text
    carga._importar()
    assert importadas == [2]
    assert len(db.listar_solicitudes(lote.id)) == 2
    assert all(s.origen == "EXCEL" for s in db.listar_solicitudes(lote.id))


def probar_reimportar_no_duplica():
    """el mismo archivo dos veces no crea solicitudes repetidas"""
    from ui.carga_masiva import CargaMasiva

    app = comun.AppFalsa()
    carga = CargaMasiva(app, lambda *_a: None)
    lote = comun.lote()
    ruta = _excel_de_prueba(comun.carpeta_temporal())
    for _ in range(2):
        carga.abrir(lote.id)
        carga._ruta = ruta
        comun.correr(carga._analizar())
        carga._importar()
    assert len(db.listar_solicitudes(lote.id)) == 2
    assert "ya existían" in app.ultimo_aviso


# --------------------------------------------------------------------------- #
#  Modal de alta desde carátulas
# --------------------------------------------------------------------------- #
def probar_alta_desde_caratulas():
    """cada carátula se vuelve una solicitud, y el nombre se puede corregir"""
    from ui.alta_caratulas import AltaDesdeCaratulas

    carpeta = comun.carpeta_temporal()
    for nombre in ("CARATULA JUAN PEREZ LOPEZ.pdf",
                   "CARATULA MARIA RUIZ SOTO.pdf",
                   "12345678.pdf"):
        comun.pdf_falso(carpeta, nombre)

    app = comun.AppFalsa()
    creadas = []
    modal = AltaDesdeCaratulas(app, lambda n: creadas.append(n))
    lote = comun.lote()
    modal.abrir(lote.id)
    modal.dd_empresa.value = "Abastecedora"
    modal.dd_sucursal.value = "Corporativo"
    # Los PDF de la prueba no son legibles: se comprueba de paso que un archivo
    # que el OCR no puede leer NO tumba la carga, y que el nombre cae al
    # respaldo del nombre del archivo.
    comun.correr(modal._cargar([carpeta]))

    assert modal.btn_importar.text == "Dar de alta 2", modal.btn_importar.text
    # El que no se reconoció se corrige en la tabla.
    sin_nombre = [b for b in modal._borradores if not b.nombre_detectado][0]
    modal._renombrar(sin_nombre, "pedro gomez soto")
    assert sin_nombre.solicitud.beneficiario_nombre == "PEDRO GOMEZ SOTO"
    assert modal.btn_importar.text == "Dar de alta 3"

    modal._importar()
    assert creadas == [3]
    con_caratula = sum(1 for s in db.listar_solicitudes(lote.id)
                       if "caratula" in documentos.de_solicitud(s.id))
    assert con_caratula == 3, "ninguna solicitud debe quedarse sin carátula"


def _pantalla_con(solicitudes: list[tuple[str, str]]):
    """Pantalla de solicitudes con un lote lleno de (nombre, tipo)."""
    from ui.solicitudes import SeccionSolicitudes

    app = comun.AppFalsa()
    pantalla = SeccionSolicitudes(app)
    pantalla.cargar_desde_db()
    for nombre, tipo in solicitudes:
        comun.solicitud(pantalla._lote.id, nombre, tipo=tipo)
    pantalla._recargar_solicitudes()
    return pantalla


def probar_el_buscador_del_listado_filtra_por_beneficiario():
    """se busca sin acentos, por palabras sueltas y en cualquier orden"""
    pantalla = _pantalla_con([
        ("JOSÉ MUÑOZ ÁVILA", "Acreedor"),
        ("ANA LOPEZ RUIZ", "Acreedor"),
        ("PEDRO GOMEZ SOTO", "Deudor"),
    ])
    assert len(pantalla._visibles()) == 3

    # Sin acentos y en minúsculas.
    pantalla.tf_buscar.value = "munoz"
    pantalla._cambiar_busqueda()
    assert [s.beneficiario_nombre for s in pantalla._visibles()] == [
        "JOSÉ MUÑOZ ÁVILA"]

    # Palabras sueltas y en cualquier orden: nadie recuerda si el apellido iba
    # antes o después.
    pantalla.tf_buscar.value = "ruiz ana"
    pantalla._cambiar_busqueda()
    assert [s.beneficiario_nombre for s in pantalla._visibles()] == [
        "ANA LOPEZ RUIZ"]

    pantalla._limpiar_busqueda()
    assert len(pantalla._visibles()) == 3


def probar_el_filtro_de_tipo_admite_uno_dos_o_los_tres():
    """el filtro de tipo de beneficiario combina cualquier subconjunto"""
    from core.catalogos import TIPOS_BENEFICIARIO

    pantalla = _pantalla_con([
        ("UNO PROVEEDOR", "Proveedor"),
        ("DOS DEUDOR", "Deudor"),
        ("TRES ACREEDOR", "Acreedor"),
    ])

    class _E:
        def __init__(self, seleccion):
            self.control = type("C", (), {"selected": seleccion})()

    # Uno solo.
    pantalla._cambiar_tipos(_E(["Proveedor"]))
    assert [s.tipo_beneficiario for s in pantalla._visibles()] == ["Proveedor"]

    # Combinación de dos.
    pantalla._cambiar_tipos(_E(["Proveedor", "Acreedor"]))
    assert sorted(s.tipo_beneficiario for s in pantalla._visibles()) == [
        "Acreedor", "Proveedor"]

    # Los tres.
    pantalla._cambiar_tipos(_E(list(TIPOS_BENEFICIARIO)))
    assert len(pantalla._visibles()) == 3


def probar_el_listado_se_pagina_y_seleccionar_todas_respeta_el_filtro():
    """con muchas filas se pagina, y «todas» marca lo filtrado, no el lote"""
    from ui.solicitudes import _POR_PAGINA

    filas = [(f"BENEFICIARIO {i:03}", "Acreedor" if i % 2 else "Proveedor")
             for i in range(_POR_PAGINA + 20)]
    pantalla = _pantalla_con(filas)

    # La primera página no trae más de lo permitido y la barra aparece.
    pantalla._pintar()
    assert pantalla.barra_paginacion.visible
    assert pantalla.btn_pag_anterior.disabled
    assert not pantalla.btn_pag_siguiente.disabled

    pantalla._mover_pagina(1)
    assert pantalla._pagina == 1
    assert not pantalla.btn_pag_anterior.disabled
    # Pasada la última no se puede seguir: la página se recorta al pintar.
    pantalla._mover_pagina(50)
    assert pantalla.btn_pag_siguiente.disabled

    # «Seleccionar todas» actúa sobre lo FILTRADO, no sobre el lote entero, y
    # tampoco solo sobre la página que se está viendo.
    class _E:
        def __init__(self, seleccion):
            self.control = type("C", (), {"selected": seleccion})()

    pantalla._cambiar_tipos(_E(["Proveedor"]))
    visibles = {s.id for s in pantalla._visibles()}
    assert 0 < len(visibles) < len(pantalla._solicitudes)

    pantalla._alternar_todos(type("E", (), {"control": type(
        "C", (), {"value": True})()})())
    assert pantalla._seleccionados == visibles, "no debe tocar lo que el filtro esconde"


def probar_la_fecha_comun_se_aplica_a_todas_las_caratulas():
    """la fecha de pago del paso 2 llena las solicitudes que no la traen"""
    # La carátula no trae fecha de pago —no es un dato del banco—, así que sin
    # esto un lote entero queda con «Falta la fecha de pago» y hay que corregir
    # solicitud por solicitud.
    from ui.alta_caratulas import AltaDesdeCaratulas

    carpeta = comun.carpeta_temporal()
    for nombre in ("CARATULA JUAN PEREZ LOPEZ.pdf",
                   "CARATULA MARIA RUIZ SOTO.pdf"):
        comun.pdf_falso(carpeta, nombre)

    app = comun.AppFalsa()
    modal = AltaDesdeCaratulas(app, lambda *_a: None)
    modal.abrir(comun.lote().id)
    comun.correr(modal._cargar([carpeta]))
    assert modal._borradores, "no se cargó ninguna carátula"
    assert all(not b.solicitud.fecha_pago for b in modal._borradores)

    modal.campo_fecha.value = "30/09/2026"
    modal._aplicar_comunes()
    assert all(b.solicitud.fecha_pago == "30/09/2026"
               for b in modal._borradores), "no se aplicó a todas"

    # Y no pisa la que ya venía puesta: ahí manda lo que trajo el Excel.
    modal._borradores[0].solicitud.fecha_pago = "15/10/2026"
    modal.campo_fecha.value = "01/12/2026"
    modal._aplicar_comunes()
    assert modal._borradores[0].solicitud.fecha_pago == "15/10/2026"


# --------------------------------------------------------------------------- #
#  Modal de asignación masiva
# --------------------------------------------------------------------------- #
def probar_el_centro_de_costos_solo_se_pide_cuando_hay_insumos():
    """un concepto de pago no lleva centro de costos; un insumo sí"""
    # Pedirlo con un alcance de puros Acreedores sería pedir un dato que no se
    # escribe en ningún lado: el renglón de concepto no tiene esos campos.
    from ui.asignacion_masiva import AsignacionMasiva

    conceptos.importar(["VIGILANCIA"], "Abastecedora")
    app = comun.AppFalsa()
    modal = AsignacionMasiva(app, lambda *_a: None)

    solo_acreedores = comun.lote()
    comun.solicitud(solo_acreedores.id, "ANA LOPEZ", tipo="Acreedor")
    modal.abrir(solo_acreedores.id, set())
    assert not modal.bloque_insumo.visible, "no hay proveedores que lo necesiten"

    con_proveedor = comun.lote()
    comun.solicitud(con_proveedor.id, "ANA LOPEZ", tipo="Acreedor")
    comun.solicitud(con_proveedor.id, "PROVEEDOR SA", tipo="Proveedor")
    modal.abrir(con_proveedor.id, set())
    assert modal.bloque_insumo.visible, "con proveedores sí hacen falta"

    # Y no deja calcular sin ellos: el renglón de insumo iría incompleto a SIPP.
    modal.dd_nombre.value = "VIGILANCIA"
    modal.tf_centro.value = ""
    modal.tf_cuenta.value = ""
    modal._calcular()
    assert modal._plan is None, "debió detenerse y pedir los datos del insumo"

    modal.tf_centro.value = "CC-01"
    modal.tf_cuenta.value = "5000"
    modal._calcular()
    assert modal._plan is not None, "con los datos completos sí procede"


def probar_la_asignacion_masiva_reasigna_la_cabecera_sin_concepto():
    """se puede fijar empresa, sucursal y fecha sin elegir concepto"""
    from ui.asignacion_masiva import AsignacionMasiva

    app = comun.AppFalsa()
    modal = AsignacionMasiva(app, lambda *_a: None)
    lote = comun.lote()
    comun.solicitud(lote.id, "ANA LOPEZ", tipo="Acreedor")
    modal.abrir(lote.id, set())

    # Sin concepto y sin cabecera no hay nada que hacer.
    modal.dd_nombre.value = ""
    modal._calcular()
    assert modal._plan is None

    modal.campo_fecha.value = "30/09/2026"
    modal._calcular()
    assert modal._plan is not None and modal._plan.cambios
    modal._aplicar()
    assert all(s.fecha_pago == "30/09/2026"
               for s in db.listar_solicitudes(lote.id))


def probar_asignacion_invalida_la_previa_al_cambiar_criterios():
    """cambiar un criterio obliga a recalcular"""
    # Si no, se aplicaría un plan distinto del que se vio, que es exactamente lo
    # que este modal existe para evitar.
    from ui.asignacion_masiva import AsignacionMasiva

    conceptos.importar(["VIGILANCIA"], "Abastecedora")
    app = comun.AppFalsa()
    modal = AsignacionMasiva(app, lambda *_a: None)
    lote = comun.lote()
    comun.solicitud(lote.id, "ANA LOPEZ", partidas=[
        comun.insumo("Facturado", 1000.0, origen="CFDI")])

    modal.abrir(lote.id, set())
    modal.dd_nombre.value = "VIGILANCIA"
    modal._calcular()
    assert modal.previa.visible and not modal.btn_aplicar.disabled

    modal.dd_modo.value = "Reemplazar los renglones existentes"
    modal._invalidar()
    assert not modal.previa.visible and modal.btn_aplicar.disabled


def probar_asignacion_respeta_la_seleccion_previa():
    """abrir con filas marcadas propone el alcance «seleccionadas»"""
    from ui.asignacion_masiva import AsignacionMasiva

    app = comun.AppFalsa()
    modal = AsignacionMasiva(app, lambda *_a: None)
    lote = comun.lote()
    s = comun.solicitud(lote.id, "ANA LOPEZ", partidas=[
        comun.insumo("Facturado", 1000.0, origen="CFDI")])
    comun.solicitud(lote.id, "OTRO", partidas=[
        comun.insumo("Facturado", 500.0, origen="CFDI")])

    # El alcance ya no se elige en el modal: lo dice la selección con la que se
    # abrió, y se enuncia para que se vea a cuántas va a aplicar.
    modal.abrir(lote.id, {s.id})
    assert modal._alcance() == asignacion.SELECCIONADAS
    assert "1 solicitud(es) que seleccionaste" in modal.txt_alcance.value
    modal.dd_nombre.value = "VIGILANCIA"
    modal._calcular()
    assert [c.solicitud.id for c in modal._plan.cambios] == [s.id]

    # Sin selección previa, el alcance es el lote entero y se advierte.
    modal.abrir(lote.id, set())
    assert modal._alcance() == asignacion.TODAS
    assert "TODAS" in modal.txt_alcance.value


def probar_asignacion_deja_aplicar_aunque_falte_otro_campo():
    """el botón no se bloquea por un campo que este modal no puede corregir"""
    # El caso reportado: una solicitud sin CLABE dejaba «Aplicar» apagado, sin
    # forma de asignar el concepto ni de arreglar la CLABE desde aquí.
    from ui.asignacion_masiva import AsignacionMasiva

    app = comun.AppFalsa()
    modal = AsignacionMasiva(app, lambda *_a: None)
    lote = comun.lote()
    s = comun.solicitud(lote.id, "ANDREA BARROS GARCIA", cuenta_clabe="",
                        partidas=[comun.insumo("Facturado", 1500.0,
                                               origen="CFDI")])
    modal.abrir(lote.id, set())
    modal.dd_nombre.value = "PAGO TARJETA CREDITO"
    modal._calcular()

    assert not modal.btn_aplicar.disabled, "tenía que poder aplicarse"
    assert modal.btn_aplicar.text == "Aplicar a 1"
    resumen = modal.txt_resumen.value
    assert "0 quedarían listas" in resumen
    assert "1 se asignan pero les falta algo fuera del desglose" in resumen
    assert "no se pueden asignar" not in resumen

    modal._aplicar()
    assert len(db.listar_partidas(s.id)) == 2
    # Y se dice que sigue sin poder capturarse: el modal se cierra y es la
    # última oportunidad de decirlo.
    assert "el robot no las capturará" in app.ultimo_aviso


def probar_asignacion_aplica_y_deshace_desde_el_aviso():
    """se aplica y se puede revertir desde el propio aviso de éxito"""
    from ui.asignacion_masiva import AsignacionMasiva

    app = comun.AppFalsa()
    modal = AsignacionMasiva(app, lambda *_a: None)
    lote = comun.lote()
    ana = comun.solicitud(lote.id, "ANA LOPEZ", partidas=[
        comun.insumo("Facturado", 1000.0, origen="CFDI")])

    modal.abrir(lote.id, set())
    modal.dd_nombre.value = "VIGILANCIA"
    modal._calcular()
    modal._aplicar()
    assert len(db.listar_partidas(ana.id)) == 2

    modal._deshacer_desde_aviso()
    assert len(db.listar_partidas(ana.id)) == 1


# --------------------------------------------------------------------------- #
#  Bitácora y conceptos
# --------------------------------------------------------------------------- #
def probar_bitacora_filtra_por_nivel():
    """el filtro de la bitácora reduce lo que se muestra"""
    from ui.bitacora import SeccionBitacora

    lote = comun.lote()
    s = comun.solicitud(lote.id, "ANA LOPEZ",
                        partidas=[comun.concepto("VIGILANCIA", 100.0)])
    db.registrar(s.id, "login", "Sesión iniciada")
    db.registrar(s.id, "guardar", "Falló", nivel="ERROR")

    app = comun.AppFalsa()
    pantalla = SeccionBitacora(app)
    pantalla.cargar_desde_db()
    assert "2 registro" in pantalla.txt_resumen.value
    pantalla._nivel = "ERROR"
    pantalla.cargar_desde_db()
    assert "1 registro" in pantalla.txt_resumen.value


def probar_bitacora_filtra_por_lote():
    """se puede mirar el historial de un lote sin el ruido de los demás"""
    from ui.bitacora import SeccionBitacora

    uno = comun.lote("Lote uno")
    otro = comun.lote("Lote dos")
    a = comun.solicitud(uno.id, "ANA LOPEZ")
    b = comun.solicitud(otro.id, "BRUNO DIAZ")
    db.registrar(a.id, "login", "Sesión iniciada")
    db.registrar(a.id, "guardar", "Guardada")
    db.registrar(b.id, "login", "Sesión iniciada")

    app = comun.AppFalsa()
    pantalla = SeccionBitacora(app)
    pantalla.cargar_desde_db()
    assert "3 registro" in pantalla.txt_resumen.value

    pantalla._lote_id = uno.id
    pantalla.cargar_desde_db()
    assert "2 registro" in pantalla.txt_resumen.value, pantalla.txt_resumen.value
    assert "1 solicitud" in pantalla.txt_resumen.value


def probar_bitacora_filtra_por_fecha():
    """el rango de fechas acota, y el día de «hasta» se incluye"""
    # El momento se guarda con hora ('2026-08-17T15:03'), así que un «hasta»
    # comparado en crudo dejaría fuera todo lo de ese mismo día.
    from ui.bitacora import SeccionBitacora

    lote = comun.lote()
    s = comun.solicitud(lote.id, "ANA LOPEZ")
    db.registrar(s.id, "login", "Sesión iniciada")

    hoy = db._ahora()[:10]
    assert len(db.listar_bitacora(desde=hoy, hasta=hoy)) == 1, (
        "el registro de hoy tiene que entrar en un rango que termina hoy")
    assert not db.listar_bitacora(desde="2000-01-01", hasta="2000-01-02")

    app = comun.AppFalsa()
    pantalla = SeccionBitacora(app)
    pantalla.campo_desde.value = "01/01/2000"
    pantalla.campo_hasta.value = "02/01/2000"
    pantalla.cargar_desde_db()
    assert "0 registro" in pantalla.txt_resumen.value
    assert pantalla.sin_resultados.visible, "hay que distinguirlo de vacía"
    assert not pantalla.vacio.visible


def probar_bitacora_agrupa_los_pasos_por_solicitud():
    """cada solicitud abre su propio bloque, en vez de una lista corrida"""
    from ui.bitacora import SeccionBitacora
    from ui.tabla_responsiva import Cabecera

    lote = comun.lote()
    a = comun.solicitud(lote.id, "ANA LOPEZ")
    b = comun.solicitud(lote.id, "BRUNO DIAZ")
    db.registrar(a.id, "login", "Sesión iniciada")
    db.registrar(a.id, "guardar", "Falló", nivel="ERROR")
    db.registrar(b.id, "login", "Sesión iniciada")

    app = comun.AppFalsa()
    pantalla = SeccionBitacora(app)
    pantalla.cargar_desde_db()

    # Dos solicitudes -> dos bandas de grupo, siempre visibles.
    cabeceras = [f for f in pantalla.tabla._filas if isinstance(f, Cabecera)]
    assert len(cabeceras) == 2, f"{len(cabeceras)} bandas"
    assert "2 solicitud(es)" in pantalla.txt_resumen.value
    assert "1 error(es)" in pantalla.txt_resumen.value

    # El bloque con error se abre solo (sus 2 pasos se ven); el otro no.
    assert a.id in pantalla._expandidos
    assert b.id not in pantalla._expandidos
    assert len(pantalla.tabla._filas) == 4, "2 bandas + los 2 pasos del fallido"


def probar_los_bloques_de_la_bitacora_se_despliegan_y_se_recuerdan():
    """cada solicitud se abre y se cierra, y el refresco no lo deshace"""
    from ui.bitacora import SeccionBitacora
    from ui.tabla_responsiva import Cabecera

    lote = comun.lote()
    s = comun.solicitud(lote.id, "ANA LOPEZ")
    db.registrar(s.id, "login", "Sesión iniciada")
    db.registrar(s.id, "guardar", "Guardada")

    app = comun.AppFalsa()
    pantalla = SeccionBitacora(app)
    pantalla.cargar_desde_db()
    # Sin errores, arranca cerrado: solo se ve su banda.
    assert len(pantalla.tabla._filas) == 1
    assert s.id not in pantalla._expandidos

    pantalla._alternar_bloque(s.id)
    assert len(pantalla.tabla._filas) == 3, "banda + sus dos pasos"

    # Un refresco NO debe cerrar lo que el usuario abrió.
    pantalla.cargar_desde_db()
    assert s.id in pantalla._expandidos
    assert len(pantalla.tabla._filas) == 3

    pantalla._todos_los_bloques(False)
    assert len(pantalla.tabla._filas) == 1
    pantalla._todos_los_bloques(True)
    assert len(pantalla.tabla._filas) == 3


def probar_un_bloque_con_error_cerrado_a_mano_no_se_reabre():
    """si lo cierras, se queda cerrado aunque tenga errores"""
    # Se abren solos la PRIMERA vez que aparecen; después manda el usuario.
    from ui.bitacora import SeccionBitacora

    lote = comun.lote()
    s = comun.solicitud(lote.id, "ANA LOPEZ")
    db.registrar(s.id, "guardar", "Falló", nivel="ERROR")

    app = comun.AppFalsa()
    pantalla = SeccionBitacora(app)
    pantalla.cargar_desde_db()
    assert s.id in pantalla._expandidos, "el que falla se abre solo"

    pantalla._alternar_bloque(s.id)
    pantalla.cargar_desde_db()
    assert s.id not in pantalla._expandidos, "no debe reabrirse en el refresco"


def probar_quitar_filtros_devuelve_la_bitacora_completa():
    """el botón de limpiar deja los cuatro filtros en su sitio"""
    from ui.bitacora import SeccionBitacora

    lote = comun.lote()
    s = comun.solicitud(lote.id, "ANA LOPEZ")
    db.registrar(s.id, "login", "Sesión iniciada")

    app = comun.AppFalsa()
    pantalla = SeccionBitacora(app)
    pantalla._nivel = "ERROR"
    pantalla._lote_id = lote.id
    pantalla.campo_desde.value = "01/01/2000"
    pantalla.campo_hasta.value = "02/01/2000"
    pantalla.cargar_desde_db()
    assert "0 registro" in pantalla.txt_resumen.value

    pantalla._quitar_filtros()
    assert not pantalla._hay_filtros()
    assert "1 registro" in pantalla.txt_resumen.value


def probar_buscar_en_el_catalogo_de_conceptos():
    """el buscador filtra sin acentos, por palabras sueltas y en cualquier orden"""
    from ui.conceptos import SeccionConceptos

    conceptos.importar(["ISR RETENCIONES POR SALARIOS", "PREVISION OBRERA",
                        "PREVISION PATRONAL", "VIGILANCIA"], "Abastecedora")
    app = comun.AppFalsa()
    pantalla = SeccionConceptos(app)
    pantalla.cargar_desde_db()
    assert "4 concepto(s)" in pantalla.txt_estado.value
    assert not pantalla.btn_limpiar.visible

    def buscar(texto):
        pantalla.tf_buscar.value = texto
        pantalla._cambiar_filtro()
        return [c.nombre for c in pantalla._conceptos
                if pantalla._coincide(c)]

    assert buscar("prevision") == ["PREVISION OBRERA", "PREVISION PATRONAL"]
    assert buscar("previsión") == ["PREVISION OBRERA", "PREVISION PATRONAL"], (
        "el acento no debe cambiar el resultado")
    assert buscar("salarios isr") == ["ISR RETENCIONES POR SALARIOS"], (
        "las palabras van sueltas y en cualquier orden")

    buscar("prevision")
    assert "2 de 4 concepto(s)" in pantalla.txt_estado.value, (
        "el total tiene que seguir a la vista al filtrar")
    assert pantalla.btn_limpiar.visible


def probar_buscar_sin_resultados_no_invita_a_importar():
    """el vacío por filtro se distingue del catálogo vacío"""
    # Con un solo mensaje, no encontrar «vigilancia» invitaría a importar de
    # SIPP un catálogo que ya está completo, o a darla de alta a mano duplicada.
    from ui.conceptos import SeccionConceptos

    app = comun.AppFalsa()
    pantalla = SeccionConceptos(app)
    pantalla.cargar_desde_db()
    assert pantalla.vacio.visible and not pantalla.sin_resultados.visible

    conceptos.importar(["VIGILANCIA"], "Abastecedora")
    pantalla.tf_buscar.value = "no existe"
    pantalla._cambiar_filtro()
    assert pantalla.sin_resultados.visible
    assert not pantalla.vacio.visible
    assert not pantalla.tabla.control.visible

    pantalla._limpiar_filtro()
    assert pantalla.tabla.control.visible
    assert not pantalla.sin_resultados.visible
    assert not pantalla.btn_limpiar.visible


def probar_el_alta_deja_ver_el_concepto_nuevo():
    """agregar con una búsqueda puesta no esconde lo recién agregado"""
    from ui.conceptos import SeccionConceptos

    conceptos.importar(["VIGILANCIA"], "Abastecedora")
    app = comun.AppFalsa()
    pantalla = SeccionConceptos(app)
    pantalla.cargar_desde_db()
    pantalla.tf_buscar.value = "vigilancia"
    pantalla._cambiar_filtro()

    pantalla._abrir_alta()
    pantalla.tf_nombre.value = "FLETES FORANEOS"
    pantalla._guardar_alta()
    assert not pantalla.tf_buscar.value, "la búsqueda debe quedar limpia"
    assert "2 concepto(s)" in pantalla.txt_estado.value


def probar_alta_manual_de_concepto_avisa_de_soporte():
    """al crear un concepto a mano se advierte que puede no existir en SIPP"""
    from ui.conceptos import SeccionConceptos

    app = comun.AppFalsa()
    pantalla = SeccionConceptos(app)
    pantalla.cargar_desde_db()
    pantalla._abrir_alta()
    pantalla.tf_nombre.value = "CONCEPTO NUEVO"
    pantalla._guardar_alta()
    assert "soporte" in app.ultimo_aviso.lower()
    assert conceptos.existe("CONCEPTO NUEVO")

    # Y el duplicado se rechaza dentro del diálogo, sin cerrarlo.
    pantalla._abrir_alta()
    pantalla.tf_nombre.value = "concepto nuevo"
    pantalla._guardar_alta()
    assert pantalla.txt_error_alta.visible


def probar_los_documentos_adjuntos_se_ven_al_guardar_sin_replegar():
    """la tabla refleja la carátula recién adjuntada, sin tocar la fila"""
    # El fallo reportado: se adjuntaba la carátula, se guardaba, y la fila
    # seguía diciendo «FALTA LA CARÁTULA» hasta plegar y desplegar el detalle.
    # La causa era el orden: se repintaba ANTES de registrar los documentos.
    from ui.solicitudes import SeccionSolicitudes

    carpeta = comun.carpeta_temporal()
    caratula = comun.pdf_falso(carpeta, "CARATULA ANA LOPEZ.pdf")

    app = comun.AppFalsa()
    pantalla = SeccionSolicitudes(app)
    pantalla.cargar_desde_db()

    s = db.Solicitud(
        lote_id=pantalla._lote.id, empresa="Abastecedora",
        sucursal="Corporativo", tipo_beneficiario="Acreedor",
        beneficiario_nombre="ANA LOPEZ", beneficiario_rfc="XAXX010101000",
        cuenta_clabe="012345678901234568", cuenta_banco="BBVA",
        forma_pago="Transferencia", tipo_gasto="No Deducible",
        fecha_pago=comun.fecha_futura(), moneda="Pesos (MXN)", estado="VALIDADA")
    partidas = [comun.concepto("PAGO PTU", 100.0)]

    # Tal como lo entrega el modal: solicitud, partidas y sus archivos juntos.
    pantalla._recibir_solicitud(s, partidas, {documentos.TIPO_CARATULA: caratula})

    guardada = db.listar_documentos(solicitud_id=s.id)
    assert guardada, "la carátula debió quedar registrada"

    # Y la tabla ya la ve: sin este orden, `de_solicitud` devolvía vacío al
    # pintar y la fila anunciaba que faltaba.
    assert documentos.de_solicitud(s.id).get("caratula"), (
        "la tabla se pinta leyendo esto; si está vacío, dirá que falta")


def probar_lo_importado_hereda_la_parada_del_lote():
    """con el lote en «Guardar y autorizar», nada nace en «Llenar y esperar»"""
    # El fallo reportado: se configuraba el lote para guardar y autorizar, pero
    # las solicitudes creadas desde caratulas o Excel nacian con la parada fija
    # del adaptador y el robot se detenia a pedir revision.
    from core.adaptadores import caratulas

    carpeta = comun.carpeta_temporal()
    comun.pdf_falso(carpeta, "CARATULA ANA LOPEZ.pdf")

    borradores = caratulas.crear_borradores([carpeta], "lote-x",
                                            parada="AUTORIZAR",
                                            leer_caratula=False)
    assert borradores, "debio crear un borrador"
    assert all(b.solicitud.parada == "AUTORIZAR" for b in borradores), (
        [b.solicitud.parada for b in borradores])


def probar_el_modal_de_caratulas_recibe_la_parada_al_abrirse():
    """la pantalla se la pasa; sin eso el adaptador usa su valor fijo"""
    from ui.alta_caratulas import AltaDesdeCaratulas

    app = comun.AppFalsa()
    modal = AltaDesdeCaratulas(app, lambda *_a: None)
    modal.abrir(comun.lote().id, parada="AUTORIZAR")
    assert modal._parada == "AUTORIZAR"


def probar_el_modal_de_excel_recibe_la_parada_al_abrirse():
    """igual para la carga masiva"""
    from ui.carga_masiva import CargaMasiva

    app = comun.AppFalsa()
    modal = CargaMasiva(app, lambda *_a: None)
    modal.abrir(comun.lote().id, parada="GUARDADA")
    assert modal._parada == "GUARDADA"


# --------------------------------------------------------------------------- #
#  El aviso de ejecución no se puede quedar colgado
# --------------------------------------------------------------------------- #
def _cuerpo_de(nombre: str) -> str:
    """El código de una función anidada de `ui.solicitudes`, como texto."""
    import inspect

    from ui import solicitudes

    fuente = inspect.getsource(solicitudes)
    ini = fuente.index(f"def {nombre}(")
    return fuente[ini:fuente.index("dialogo.update()", ini)]


def probar_detener_despierta_las_dos_esperas_del_motor():
    """el motor duerme en dos sitios y Detener tiene que valer en los dos"""
    # Reportado: al terminar el lote, el aviso se quedaba en pantalla sin
    # avanzar y el reporte final no llegaba. El motor estaba dormido esperando
    # a que se cerrara el navegador, y Detener solo despertaba la otra espera
    # —la de la revisión—, así que el hilo no volvía nunca.
    cuerpo = _cuerpo_de("cancelar")
    assert "pausa.set()" in cuerpo, "la espera de la revisión"
    assert "cierre.set()" in cuerpo, (
        "y la del cierre del navegador, o el aviso se queda colgado")


def probar_cerrar_el_navegador_acusa_recibo():
    """matar Chromium tarda: sin acuse, parece que el botón no hizo nada"""
    cuerpo = _cuerpo_de("cerrar_navegador")
    assert "Cerrando" in cuerpo
    assert "disabled = True" in cuerpo, "y no se puede pulsar dos veces"
